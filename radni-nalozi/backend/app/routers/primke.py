"""Primke robe — skladištar fotografira otpremnicu dobavljača (1+ stranica), AI izvuče stavke."""
import os
from datetime import datetime
from io import BytesIO

from fastapi import APIRouter, Body, Depends, File, HTTPException, Response, UploadFile
from sqlalchemy.orm import Session

from ..ai import ai_dostupan, parsiraj_otpremnicu
from ..auth import zahtijevaj_uloge
from ..config import settings
from ..database import get_db
from ..models import Korisnik, Primka, Uloga
from ..schemas import (
    KorisnikOut,
    OtpremnicaProcitana,
    OtpremnicaZahtjev,
    PrimkaCreate,
    PrimkaOut,
    PrimkaStavka,
    PrimkaUpdate,
    RazduzenjeIn,
    StranicaOtpremnice,
)
from ..storage import spremi_sliku

router = APIRouter(prefix="/primke", tags=["primke"])

# Skladište: skladištari unose primke; voditelj/poslovođa ih vide i uređuju.
skladiste = zahtijevaj_uloge(Uloga.skladistar, Uloga.voditelj, Uloga.poslovodja)

_MEDIJ = {".jpg": "image/jpeg", ".jpeg": "image/jpeg", ".png": "image/png",
          ".webp": "image/webp", ".gif": "image/gif"}


def _ucitaj_sliku(web_putanja: str) -> tuple[bytes, str] | None:
    """Web putanja '/uploads/xxx.jpg' → (bajtovi, media_type) s diska, ili None."""
    naziv = os.path.basename(web_putanja or "")
    disk = os.path.join(settings.upload_dir, naziv)
    if not naziv or not os.path.isfile(disk):
        return None
    ext = os.path.splitext(naziv)[1].lower()
    with open(disk, "rb") as f:
        return f.read(), _MEDIJ.get(ext, "image/jpeg")


def _parse_datum(s):
    if not s:
        return None
    t = str(s).strip()
    for fmt in ("%d.%m.%Y", "%d.%m.%Y.", "%Y-%m-%d", "%d/%m/%Y", "%d. %m. %Y"):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            continue
    return None


def _stavke(sirove) -> list[PrimkaStavka]:
    out: list[PrimkaStavka] = []
    for s in (sirove or []):
        artikl = str(s.get("artikl") or "").strip()
        if not artikl:
            continue
        try:
            kol = float(s.get("kolicina")) if s.get("kolicina") not in (None, "") else 1
        except (TypeError, ValueError):
            kol = 1
        jed = str(s.get("jedinica") or "kom").strip() or "kom"
        try:
            cij = float(s.get("cijena")) if s.get("cijena") not in (None, "") else None
        except (TypeError, ValueError):
            cij = None
        out.append(PrimkaStavka(artikl=artikl, kolicina=kol or 1, jedinica=jed, cijena=cij))
    return out


@router.post("/stranica", response_model=StranicaOtpremnice)
def stranica(slika: UploadFile = File(...), _: Korisnik = Depends(skladiste)):
    """Spremi jednu stranicu otpremnice (bez AI-ja). Vrati web putanju."""
    return StranicaOtpremnice(foto=spremi_sliku(slika))


@router.post("/procitaj", response_model=OtpremnicaProcitana)
def procitaj(zahtjev: OtpremnicaZahtjev, _: Korisnik = Depends(skladiste)):
    """AI pročita sve priložene stranice otpremnice (uzastopne stranice jednog dokumenta)."""
    if not ai_dostupan():
        return OtpremnicaProcitana(ai=False)
    slike = [s for s in (_ucitaj_sliku(p) for p in zahtjev.fotke) if s]
    if not slike:
        raise HTTPException(status_code=400, detail="Nema slika za čitanje.")
    try:
        podaci = parsiraj_otpremnicu(slike)
    except Exception as e:  # noqa: BLE001
        raise HTTPException(status_code=502, detail=f"AI nije uspio pročitati otpremnicu: {e}")
    return OtpremnicaProcitana(
        dobavljac=(podaci.get("dobavljac") or None),
        broj_otpremnice=(podaci.get("broj_otpremnice") or None),
        datum=_parse_datum(podaci.get("datum")),
        stavke=_stavke(podaci.get("stavke")),
        ai=True,
    )


@router.get("", response_model=list[PrimkaOut])
def popis(_: Korisnik = Depends(skladiste), db: Session = Depends(get_db)):
    return db.query(Primka).order_by(Primka.kreiran.desc()).all()


@router.get("/narucitelji", response_model=list[KorisnikOut])
def narucitelji(_: Korisnik = Depends(skladiste), db: Session = Depends(get_db)):
    """Korisnici označeni kao naručitelji (kvačica u šifrarniku) — za dropdown."""
    return (
        db.query(Korisnik)
        .filter(Korisnik.aktivan.is_(True), Korisnik.narucitelj.is_(True))
        .order_by(Korisnik.ime)
        .all()
    )


def _kol_cij(s: dict):
    try:
        kol = float(s.get("kolicina")) if s.get("kolicina") not in (None, "") else None
    except (TypeError, ValueError):
        kol = None
    try:
        cij = float(s.get("cijena")) if s.get("cijena") not in (None, "") else None
    except (TypeError, ValueError):
        cij = None
    iznos = round((kol or 0) * cij, 2) if cij is not None else None
    return kol, cij, iznos


def _xlsx_odgovor(wb, naziv: str) -> Response:
    sigurno = "".join(c if (c.isalnum() or c in " -_.") else "_" for c in (naziv or "primka")) or "primka"
    buf = BytesIO()
    wb.save(buf)
    return Response(
        content=buf.getvalue(),
        media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        headers={"Content-Disposition": f'attachment; filename="{sigurno}.xlsx"'},
    )


@router.get("/izvoz.xlsx")
def izvoz(_: Korisnik = Depends(skladiste), db: Session = Depends(get_db)):
    """Izvoz SVIH primki u Excel — jedan red po stavci (ravna tablica za analizu)."""
    from openpyxl import Workbook

    wb = Workbook()
    ws = wb.active
    ws.title = "Primke"
    ws.append(["Dobavljač", "Broj otpremnice", "Datum", "Artikl",
               "Količina", "Jedinica", "Cijena (EUR)", "Iznos (EUR)"])
    for c in ws[1]:
        c.font = c.font.copy(bold=True)
    for p in db.query(Primka).order_by(Primka.kreiran.desc()).all():
        stavke = p.stavke or []
        if not stavke:
            ws.append([p.dobavljac, p.broj_otpremnice, p.datum, None, None, None, None, None])
            continue
        for s in stavke:
            kol, cij, iznos = _kol_cij(s)
            ws.append([p.dobavljac, p.broj_otpremnice, p.datum, s.get("artikl"),
                       kol, s.get("jedinica"), cij, iznos])
    for red in ws.iter_rows(min_row=2, min_col=3, max_col=3):
        red[0].number_format = "DD.MM.YYYY"
    for stupac, sirina in zip("ABCDEFGH", (26, 18, 12, 44, 10, 10, 13, 13)):
        ws.column_dimensions[stupac].width = sirina
    return _xlsx_odgovor(wb, "primke")


@router.get("/{primka_id}", response_model=PrimkaOut)
def dohvati(primka_id: int, _: Korisnik = Depends(skladiste), db: Session = Depends(get_db)):
    p = db.get(Primka, primka_id)
    if not p:
        raise HTTPException(status_code=404, detail="Primka ne postoji")
    return p


@router.get("/{primka_id}/izvoz.xlsx")
def izvoz_jedne(primka_id: int, _: Korisnik = Depends(skladiste), db: Session = Depends(get_db)):
    """Izvoz JEDNE primke u Excel (zaglavlje + stavke s iznosom + ukupno)."""
    from openpyxl import Workbook

    p = db.get(Primka, primka_id)
    if not p:
        raise HTTPException(status_code=404, detail="Primka ne postoji")

    wb = Workbook()
    ws = wb.active
    ws.title = "Primka"
    ws.append(["Dobavljač", p.dobavljac or ""])
    ws.append(["Broj otpremnice", p.broj_otpremnice or ""])
    ws.append(["Datum", p.datum])
    if p.napomena:
        ws.append(["Napomena", p.napomena])
    ws["B3"].number_format = "DD.MM.YYYY"
    ws.append([])
    red_zaglavlja = ws.max_row + 1
    ws.append(["Artikl", "Količina", "Jedinica", "Cijena (EUR)", "Iznos (EUR)"])
    for c in ws[red_zaglavlja]:
        c.font = c.font.copy(bold=True)

    ukupno = 0.0
    ima = False
    for s in (p.stavke or []):
        kol, cij, iznos = _kol_cij(s)
        ws.append([s.get("artikl"), kol, s.get("jedinica"), cij, iznos])
        if iznos is not None:
            ukupno += iznos
            ima = True
    if ima:
        ws.append(["", "", "", "Ukupno", round(ukupno, 2)])
        for c in ws[ws.max_row]:
            c.font = c.font.copy(bold=True)

    for stupac, sirina in zip("ABCDE", (46, 10, 10, 13, 13)):
        ws.column_dimensions[stupac].width = sirina
    return _xlsx_odgovor(wb, f"primka-{p.broj_otpremnice or p.id}")


@router.post("", response_model=PrimkaOut, status_code=201)
def kreiraj(podaci: PrimkaCreate, korisnik: Korisnik = Depends(skladiste), db: Session = Depends(get_db)):
    stavke = [x.model_dump() for x in podaci.stavke if x.artikl.strip()]
    fotografije = [f for f in (podaci.fotografije or []) if f]
    if not fotografije and podaci.foto:            # kompatibilnost sa starim klijentom
        fotografije = [podaci.foto]
    if not (podaci.dobavljac or podaci.broj_otpremnice or stavke or fotografije):
        raise HTTPException(status_code=400, detail="Primka je prazna — dodaj barem dobavljača, broj otpremnice ili stavku.")
    p = Primka(
        dobavljac=(podaci.dobavljac or None),
        broj_otpremnice=(podaci.broj_otpremnice or None),
        datum=podaci.datum,
        napomena=(podaci.napomena or None),
        foto_putanja=(fotografije[0] if fotografije else None),
        fotografije=fotografije or None,
        stavke=stavke,
        kreirao_id=korisnik.id,
    )
    db.add(p)
    db.commit()
    db.refresh(p)
    return p


@router.patch("/{primka_id}", response_model=PrimkaOut)
def azuriraj(primka_id: int, podaci: PrimkaUpdate, _: Korisnik = Depends(skladiste), db: Session = Depends(get_db)):
    p = db.get(Primka, primka_id)
    if not p:
        raise HTTPException(status_code=404, detail="Primka ne postoji")
    izmjene = podaci.model_dump(exclude_unset=True)
    if izmjene.get("stavke") is not None:
        izmjene["stavke"] = [s for s in izmjene["stavke"] if (s.get("artikl") or "").strip()]
    for polje, vrijednost in izmjene.items():
        setattr(p, polje, vrijednost)
    db.commit()
    db.refresh(p)
    return p


@router.delete("/{primka_id}", status_code=204)
def obrisi(primka_id: int, _: Korisnik = Depends(skladiste), db: Session = Depends(get_db)):
    p = db.get(Primka, primka_id)
    if not p:
        raise HTTPException(status_code=404, detail="Primka ne postoji")
    db.delete(p)
    db.commit()


# --- Dashboard skladišta: naručitelj po dijelu + razduženje na vozila ---------
def _stavke_kopija(p: Primka, idx: int) -> list[dict]:
    stavke = [dict(s) for s in (p.stavke or [])]
    if idx < 0 or idx >= len(stavke):
        raise HTTPException(status_code=404, detail="Stavka ne postoji")
    return stavke


@router.patch("/{primka_id}/stavka/{idx}", response_model=PrimkaOut)
def uredi_stavku(primka_id: int, idx: int, narucio_id: int | None = Body(default=None, embed=True),
                 _: Korisnik = Depends(skladiste), db: Session = Depends(get_db)):
    """Postavi poslovođu koji je naručio taj dio (po stavci)."""
    p = db.get(Primka, primka_id)
    if not p:
        raise HTTPException(status_code=404, detail="Primka ne postoji")
    stavke = _stavke_kopija(p, idx)
    stavke[idx] = {**stavke[idx], "narucio_id": narucio_id}
    p.stavke = stavke
    db.commit()
    db.refresh(p)
    return p


@router.post("/{primka_id}/stavka/{idx}/razduzenje", response_model=PrimkaOut)
def dodaj_razduzenje(primka_id: int, idx: int, podaci: RazduzenjeIn,
                     _: Korisnik = Depends(skladiste), db: Session = Depends(get_db)):
    """Dio je otišao na vozilo (GB) u nekoj količini."""
    gb = (podaci.gb or "").strip()
    if not gb:
        raise HTTPException(status_code=400, detail="Upiši garažni broj (GB).")
    try:
        kol = float(podaci.kolicina)
    except (TypeError, ValueError):
        kol = 0
    if kol <= 0:
        raise HTTPException(status_code=400, detail="Količina mora biti veća od 0.")
    p = db.get(Primka, primka_id)
    if not p:
        raise HTTPException(status_code=404, detail="Primka ne postoji")
    stavke = _stavke_kopija(p, idx)
    razd = list(stavke[idx].get("razduzenja") or [])
    razd.append({"gb": gb, "kolicina": kol})
    stavke[idx] = {**stavke[idx], "razduzenja": razd}
    p.stavke = stavke
    db.commit()
    db.refresh(p)
    return p


@router.delete("/{primka_id}/stavka/{idx}/razduzenje/{rid}", response_model=PrimkaOut)
def obrisi_razduzenje(primka_id: int, idx: int, rid: int,
                      _: Korisnik = Depends(skladiste), db: Session = Depends(get_db)):
    p = db.get(Primka, primka_id)
    if not p:
        raise HTTPException(status_code=404, detail="Primka ne postoji")
    stavke = _stavke_kopija(p, idx)
    razd = list(stavke[idx].get("razduzenja") or [])
    if rid < 0 or rid >= len(razd):
        raise HTTPException(status_code=404, detail="Razduženje ne postoji")
    razd.pop(rid)
    stavke[idx] = {**stavke[idx], "razduzenja": razd}
    p.stavke = stavke
    db.commit()
    db.refresh(p)
    return p
